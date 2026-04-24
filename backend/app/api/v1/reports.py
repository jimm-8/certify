from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime, timedelta
from io import BytesIO
from typing import Optional

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.chart import (
    AreaChart,
    BarChart,
    LineChart,
    PieChart,
    Reference,
    ScatterChart,
    Series,
)
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.database import get_db
from app.models.certificate_request import RequestStatus
from app.api.v1.auth import require_permissions
from app.repositories import CertificateRequestRepository
from app.services.release_hold_service import (
    get_effective_processing_seconds,
    get_request_age_days,
)

router = APIRouter(prefix="/reports", tags=["Reports"])


def _safe_date(value):
    if not value:
        return None
    if hasattr(value, "date"):
        return value.date()
    return value


def _excel_safe_datetime(value):
    if not value:
        return None
    if isinstance(value, datetime):
        if value.tzinfo is not None:
            return value.replace(tzinfo=None)
        return value
    return value


def _apply_period(requests, period: Optional[str], today):
    if not period or period == "all":
        return requests
    if period == "today":
        return [r for r in requests if _safe_date(r.created_at) == today]
    if period == "last_7_days":
        start = today - timedelta(days=6)
        return [
            r
            for r in requests
            if _safe_date(r.created_at) and _safe_date(r.created_at) >= start
        ]
    if period == "last_30_days":
        start = today - timedelta(days=29)
        return [
            r
            for r in requests
            if _safe_date(r.created_at) and _safe_date(r.created_at) >= start
        ]
    if period == "this_month":
        start = today.replace(day=1)
        return [
            r
            for r in requests
            if _safe_date(r.created_at) and _safe_date(r.created_at) >= start
        ]
    if period == "this_year":
        start = today.replace(month=1, day=1)
        return [
            r
            for r in requests
            if _safe_date(r.created_at) and _safe_date(r.created_at) >= start
        ]
    return requests


def _format_duration(seconds: float | None) -> str:
    if seconds is None:
        return "0"
    total_minutes = int(round(seconds / 60))
    if total_minutes <= 0:
        return "0m"
    days, rem = divmod(total_minutes, 60 * 24)
    hours, minutes = divmod(rem, 60)
    if days > 0:
        return f"{days}d {hours}h"
    if hours > 0:
        return f"{hours}h {minutes}m"
    return f"{minutes}m"


def _format_generation_duration(milliseconds: float | None) -> str:
    if milliseconds is None:
        return "0 ms"
    if milliseconds < 1000:
        return f"{int(round(milliseconds))} ms"
    seconds = milliseconds / 1000
    if seconds < 60:
        return f"{seconds:.2f} s"
    minutes = int(seconds // 60)
    remaining_seconds = seconds % 60
    return f"{minutes}m {remaining_seconds:.1f}s"


def _last_n_days(n: int, today):
    return [today - timedelta(days=i) for i in range(n - 1, -1, -1)]


def _linear_forecast(counts: list[int], horizon: int) -> list[int]:
    n = len(counts)
    if n == 0:
        return [0 for _ in range(horizon)]
    x_vals = list(range(n))
    x_mean = sum(x_vals) / n
    y_mean = sum(counts) / n
    denom = sum((x - x_mean) ** 2 for x in x_vals)
    if denom == 0:
        slope = 0
    else:
        slope = sum((x - x_mean) * (y - y_mean) for x, y in zip(x_vals, counts)) / denom
    intercept = y_mean - slope * x_mean
    preds = []
    for i in range(n, n + horizon):
        val = intercept + slope * i
        preds.append(max(0, int(round(val))))
    return preds


def _moving_average_forecast(
    counts: list[int], horizon: int, window: int = 7
) -> list[int]:
    if not counts:
        return [0 for _ in range(horizon)]
    window = min(window, len(counts))
    avg = sum(counts[-window:]) / window
    return [max(0, int(round(avg))) for _ in range(horizon)]


def _exp_smoothing_forecast(
    counts: list[int], horizon: int, alpha: float = 0.3
) -> list[int]:
    if not counts:
        return [0 for _ in range(horizon)]
    level = counts[0]
    for value in counts[1:]:
        level = alpha * value + (1 - alpha) * level
    return [max(0, int(round(level))) for _ in range(horizon)]


def _build_reports_payload(requests, period: Optional[str], now: datetime):
    today = now.date()
    filtered = _apply_period(requests, period, today)

    total_requests = len(filtered)
    released_requests = [r for r in filtered if r.status == RequestStatus.RELEASED]
    rejected_requests = [r for r in filtered if r.status == RequestStatus.REJECTED]

    pending_count = sum(
        1 for r in filtered if r.status in {RequestStatus.SUBMITTED, RequestStatus.PENDING}
    )
    processing_count = sum(
        1 for r in filtered if r.status == RequestStatus.PROCESSING
    )
    requests_today = sum(
        1 for r in filtered if _safe_date(r.created_at) == today
    )

    release_rate = (
        round((len(released_requests) / total_requests) * 100)
        if total_requests
        else 0
    )
    rejection_rate = (
        round((len(rejected_requests) / total_requests) * 100)
        if total_requests
        else 0
    )

    status_counts = Counter(
        r.status.value if hasattr(r.status, "value") else str(r.status)
        for r in filtered
        if r.status
    )

    avg_processing_seconds = None
    durations = []
    for r in released_requests:
        effective_seconds = get_effective_processing_seconds(r)
        if effective_seconds is not None:
            durations.append(effective_seconds)
    if durations:
        avg_processing_seconds = sum(durations) / len(durations)

    pdf_generation_durations = [
        r.pdf_generation_time_ms
        for r in filtered
        if getattr(r, "pdf_generation_time_ms", None) is not None
    ]
    avg_pdf_generation_ms = (
        sum(pdf_generation_durations) / len(pdf_generation_durations)
        if pdf_generation_durations
        else None
    )

    certificate_counter = Counter(
        r.certificate_type_name for r in filtered if r.certificate_type_name
    )
    program_counter = Counter(r.program for r in filtered if r.program)
    requestor_counter = Counter(r.requestor_name for r in filtered if r.requestor_name)

    last_30_days = _last_n_days(30, today)
    daily_counts = {d: 0 for d in last_30_days}
    for r in requests:
        d = _safe_date(r.created_at)
        if d in daily_counts:
            daily_counts[d] += 1

    daily_series = [
        {"date": d.isoformat(), "count": daily_counts[d]} for d in last_30_days
    ]

    avg_daily_requests = (
        round(sum(daily_counts.values()) / len(daily_counts), 2) if daily_counts else 0
    )

    aging_buckets = {
        "0-2_days": 0,
        "3-5_days": 0,
        "6-10_days": 0,
        "11+_days": 0,
    }
    bottleneck_by_status = Counter()
    bottleneck_by_certificate = defaultdict(list)
    bottleneck_by_program = defaultdict(list)
    stalled_requests = []

    for r in filtered:
        if r.status == RequestStatus.RELEASED:
            continue
        created = _safe_date(r.created_at)
        if not created:
            continue
        age_days = get_request_age_days(r, now=now)
        if age_days <= 2:
            aging_buckets["0-2_days"] += 1
        elif age_days <= 5:
            aging_buckets["3-5_days"] += 1
        elif age_days <= 10:
            aging_buckets["6-10_days"] += 1
        else:
            aging_buckets["11+_days"] += 1

        if age_days >= 5:
            status_label = (
                r.status.value if hasattr(r.status, "value") else str(r.status)
            )
            bottleneck_by_status[status_label] += 1
            if r.certificate_type_name:
                bottleneck_by_certificate[r.certificate_type_name].append(age_days)
            if r.program:
                bottleneck_by_program[r.program].append(age_days)
        if age_days >= 7:
            stalled_requests.append(
                {
                    "reference_number": r.reference_number,
                    "student_name": r.student_name,
                    "status": (
                        r.status.value if hasattr(r.status, "value") else str(r.status)
                    ),
                    "age_days": age_days,
                }
            )

    stalled_requests = sorted(
        stalled_requests, key=lambda x: x["age_days"], reverse=True
    )[:8]

    bottleneck_cert_list = [
        {
            "certificate_type": name,
            "avg_age_days": round(sum(ages) / len(ages), 1),
            "count": len(ages),
        }
        for name, ages in bottleneck_by_certificate.items()
    ]
    bottleneck_cert_list.sort(
        key=lambda x: (x["avg_age_days"], x["count"]), reverse=True
    )

    bottleneck_program_list = [
        {
            "program": name,
            "avg_age_days": round(sum(ages) / len(ages), 1),
            "count": len(ages),
        }
        for name, ages in bottleneck_by_program.items()
    ]
    bottleneck_program_list.sort(
        key=lambda x: (x["avg_age_days"], x["count"]), reverse=True
    )

    counts_for_forecast = [d["count"] for d in daily_series]
    forecast_horizon = 7
    linear_pred = _linear_forecast(counts_for_forecast, forecast_horizon)
    avg_pred = _moving_average_forecast(counts_for_forecast, forecast_horizon, window=7)
    exp_pred = _exp_smoothing_forecast(counts_for_forecast, forecast_horizon, alpha=0.3)

    forecast_days = _last_n_days(
        forecast_horizon, today + timedelta(days=forecast_horizon)
    )
    forecast = []
    for i in range(forecast_horizon):
        forecast.append(
            {
                "date": forecast_days[i].isoformat(),
                "linear": linear_pred[i],
                "moving_avg": avg_pred[i],
                "exp_smoothing": exp_pred[i],
                "blended": int(round((linear_pred[i] + avg_pred[i] + exp_pred[i]) / 3)),
            }
        )

    return {
        "generated_at": now.isoformat(),
        "period": period or "all",
        "descriptive": {
            "total_requests": total_requests,
            "released_total": len(released_requests),
            "rejected_total": len(rejected_requests),
            "release_rate": release_rate,
            "rejection_rate": rejection_rate,
            "avg_processing_time_label": _format_duration(avg_processing_seconds),
            "avg_pdf_generation_time_ms": (
                round(avg_pdf_generation_ms, 2)
                if avg_pdf_generation_ms is not None
                else None
            ),
            "avg_pdf_generation_time_label": _format_generation_duration(
                avg_pdf_generation_ms
            ),
            "pdf_generation_samples": len(pdf_generation_durations),
            "avg_daily_requests_last_30_days": avg_daily_requests,
            "status_breakdown": status_counts,
            "pending": pending_count,
            "processing": processing_count,
            "requests_today": requests_today,
            "sla_days": 2,
            "certificate_types": [
                {"name": name, "count": count}
                for name, count in certificate_counter.most_common(8)
            ],
            "programs": [
                {"name": name, "count": count}
                for name, count in program_counter.most_common(8)
            ],
            "top_requestors": [
                {"name": name, "count": count}
                for name, count in requestor_counter.most_common(5)
            ],
            "daily_requests": daily_series,
        },
        "diagnostic": {
            "aging_buckets": aging_buckets,
            "bottleneck_by_status": bottleneck_by_status,
            "bottleneck_by_certificate": bottleneck_cert_list[:6],
            "bottleneck_by_program": bottleneck_program_list[:6],
            "stalled_requests": stalled_requests,
        },
        "predictive": {
            "methodology": {
                "linear_trend_window_days": len(counts_for_forecast),
                "moving_average_window_days": min(7, len(counts_for_forecast)),
                "exp_smoothing_alpha": 0.3,
            },
            "forecast_next_7_days": forecast,
        },
    }


def _style_sheet_title(ws, title: str, subtitle: Optional[str] = None):
    ws["A1"] = title
    ws["A1"].font = Font(size=15, bold=True)
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = Font(size=10, italic=True)


def _write_table(ws, start_row: int, start_col: int, headers: list[str], rows: list[list]):
    for col_offset, header in enumerate(headers):
        cell = ws.cell(row=start_row, column=start_col + col_offset, value=header)
        cell.font = Font(bold=True)

    for row_offset, row in enumerate(rows, start=1):
        for col_offset, value in enumerate(row):
            ws.cell(row=start_row + row_offset, column=start_col + col_offset, value=value)

    return start_row + len(rows)


def _build_summary_workbook(summary: dict):
    wb = Workbook()
    overview = wb.active
    overview.title = "Overview"
    _style_sheet_title(
        overview,
        "Certify Summary Report",
        f"Generated at {summary['generated_at']} | Period: {summary['period']}",
    )

    descriptive = summary["descriptive"]
    diagnostic = summary["diagnostic"]
    predictive = summary["predictive"]

    overview_rows = [
        ["Total Requests", descriptive["total_requests"]],
        ["Released", descriptive["released_total"]],
        ["Rejected", descriptive["rejected_total"]],
        ["Release Rate (%)", descriptive["release_rate"]],
        ["Rejection Rate (%)", descriptive["rejection_rate"]],
        ["Requests Today", descriptive["requests_today"]],
        ["Pending", descriptive["pending"]],
        ["Processing", descriptive["processing"]],
        ["Average Processing Time", descriptive["avg_processing_time_label"]],
        ["Average PDF Generation", descriptive["avg_pdf_generation_time_label"]],
        ["PDF Samples", descriptive["pdf_generation_samples"]],
        ["Average Daily Requests (30d)", descriptive["avg_daily_requests_last_30_days"]],
    ]
    _write_table(overview, 4, 1, ["Metric", "Value"], overview_rows)
    status_items = list(descriptive["status_breakdown"].items())
    _write_table(
        overview,
        4,
        4,
        ["Status", "Count"],
        [[status, count] for status, count in status_items],
    )
    cert_rows = [[item["name"], item["count"]] for item in descriptive["certificate_types"]]
    overview["A19"] = "Top Certificate Types"
    overview["A19"].font = Font(size=12, bold=True)
    cert_end_row = _write_table(
        overview, 20, 1, ["Certificate Type", "Requests"], cert_rows
    )
    program_rows = [[item["name"], item["count"]] for item in descriptive["programs"]]
    overview["J19"] = "Top Programs"
    overview["J19"].font = Font(size=12, bold=True)
    program_end_row = _write_table(
        overview, 20, 10, ["Program", "Requests"], program_rows
    )
    overview.column_dimensions["A"].width = 28
    overview.column_dimensions["B"].width = 20
    overview.column_dimensions["D"].width = 20
    overview.column_dimensions["E"].width = 12
    overview.column_dimensions["J"].width = 18
    overview.column_dimensions["K"].width = 14

    if status_items:
        pie = PieChart()
        pie.title = "Status Breakdown"
        labels = Reference(overview, min_col=4, min_row=5, max_row=4 + len(status_items))
        data = Reference(overview, min_col=5, min_row=4, max_row=4 + len(status_items))
        pie.add_data(data, titles_from_data=True)
        pie.set_categories(labels)
        pie.height = 8
        pie.width = 10
        pie.dataLabels = DataLabelList()
        pie.dataLabels.showPercent = True
        overview.add_chart(pie, "G4")

    if cert_rows:
        cert_chart = BarChart()
        cert_chart.type = "bar"
        cert_chart.style = 10
        cert_chart.title = "Top Certificate Types"
        cert_chart.y_axis.title = "Certificate Type"
        cert_chart.x_axis.title = "Requests"
        cert_data = Reference(overview, min_col=2, min_row=20, max_row=cert_end_row)
        cert_labels = Reference(overview, min_col=1, min_row=21, max_row=cert_end_row)
        cert_chart.add_data(cert_data, titles_from_data=True)
        cert_chart.set_categories(cert_labels)
        cert_chart.height = 8
        cert_chart.width = 14
        overview.add_chart(cert_chart, "D19")

    if program_rows:
        program_chart = BarChart()
        program_chart.type = "bar"
        program_chart.style = 11
        program_chart.title = "Top Programs"
        program_chart.y_axis.title = "Program"
        program_chart.x_axis.title = "Requests"
        program_data = Reference(overview, min_col=11, min_row=20, max_row=program_end_row)
        program_labels = Reference(overview, min_col=10, min_row=21, max_row=program_end_row)
        program_chart.add_data(program_data, titles_from_data=True)
        program_chart.set_categories(program_labels)
        program_chart.height = 8
        program_chart.width = 14
        overview.add_chart(program_chart, "M19")

    descriptive_ws = wb.create_sheet("Descriptive")
    _style_sheet_title(
        descriptive_ws,
        "Descriptive Analytics",
        f"Generated at {summary['generated_at']} | Period: {summary['period']}",
    )

    status_end_row = _write_table(
        descriptive_ws,
        4,
        1,
        ["Status", "Count"],
        [[status, count] for status, count in status_items],
    )
    descriptive_ws["D4"] = "Top Certificate Types"
    descriptive_ws["D4"].font = Font(size=12, bold=True)
    cert_end_row = _write_table(
        descriptive_ws, 5, 4, ["Certificate Type", "Requests"], cert_rows
    )
    descriptive_ws["G4"] = "Top Programs"
    descriptive_ws["G4"].font = Font(size=12, bold=True)
    program_end_row = _write_table(
        descriptive_ws, 5, 7, ["Program", "Requests"], program_rows
    )
    descriptive_ws["J4"] = "Top Requestors"
    descriptive_ws["J4"].font = Font(size=12, bold=True)
    requestor_rows = [[item["name"], item["count"]] for item in descriptive["top_requestors"]]
    _write_table(
        descriptive_ws, 5, 10, ["Requestor", "Requests"], requestor_rows
    )

    if status_items:
        pie = PieChart()
        pie.title = "Status Breakdown"
        labels = Reference(descriptive_ws, min_col=1, min_row=5, max_row=status_end_row)
        data = Reference(descriptive_ws, min_col=2, min_row=4, max_row=status_end_row)
        pie.add_data(data, titles_from_data=True)
        pie.set_categories(labels)
        pie.height = 8
        pie.width = 10
        pie.dataLabels = DataLabelList()
        pie.dataLabels.showPercent = True
        descriptive_ws.add_chart(pie, "A20")

    if cert_rows:
        cert_chart = BarChart()
        cert_chart.type = "bar"
        cert_chart.style = 10
        cert_chart.title = "Top Certificate Types"
        cert_chart.y_axis.title = "Certificate Type"
        cert_chart.x_axis.title = "Requests"
        cert_data = Reference(descriptive_ws, min_col=5, min_row=5, max_row=cert_end_row)
        cert_labels = Reference(descriptive_ws, min_col=4, min_row=6, max_row=cert_end_row)
        cert_chart.add_data(cert_data, titles_from_data=True)
        cert_chart.set_categories(cert_labels)
        cert_chart.height = 8
        cert_chart.width = 14
        descriptive_ws.add_chart(cert_chart, "G20")

    if program_rows:
        program_chart = BarChart()
        program_chart.type = "bar"
        program_chart.style = 11
        program_chart.title = "Top Programs"
        program_chart.y_axis.title = "Program"
        program_chart.x_axis.title = "Requests"
        program_data = Reference(descriptive_ws, min_col=8, min_row=5, max_row=program_end_row)
        program_labels = Reference(descriptive_ws, min_col=7, min_row=6, max_row=program_end_row)
        program_chart.add_data(program_data, titles_from_data=True)
        program_chart.set_categories(program_labels)
        program_chart.height = 8
        program_chart.width = 14
        descriptive_ws.add_chart(program_chart, "N20")

    diagnostic_ws = wb.create_sheet("Diagnostic")
    _style_sheet_title(
        diagnostic_ws,
        "Diagnostic Analytics",
        f"Generated at {summary['generated_at']} | Period: {summary['period']}",
    )
    aging_rows = [[label.replace("_", " "), value] for label, value in diagnostic["aging_buckets"].items()]
    aging_end_row = _write_table(diagnostic_ws, 4, 1, ["Aging Bucket", "Count"], aging_rows)
    bottleneck_rows = [[label, value] for label, value in diagnostic["bottleneck_by_status"].items()]
    bottleneck_end_row = _write_table(diagnostic_ws, 4, 4, ["Status", "Count"], bottleneck_rows)

    if aging_rows:
        aging_chart = AreaChart()
        aging_chart.title = "Aging Buckets"
        aging_chart.style = 13
        aging_chart.y_axis.title = "Requests"
        aging_chart.x_axis.title = "Bucket"
        aging_data = Reference(diagnostic_ws, min_col=2, min_row=4, max_row=aging_end_row)
        aging_labels = Reference(diagnostic_ws, min_col=1, min_row=5, max_row=aging_end_row)
        aging_chart.add_data(aging_data, titles_from_data=True)
        aging_chart.set_categories(aging_labels)
        aging_chart.height = 7
        aging_chart.width = 11
        diagnostic_ws.add_chart(aging_chart, "G4")

    if bottleneck_rows:
        bottleneck_chart = BarChart()
        bottleneck_chart.title = "Bottlenecks by Status"
        bottleneck_chart.style = 12
        bottleneck_chart.y_axis.title = "Requests"
        bottleneck_chart.x_axis.title = "Status"
        bottleneck_data = Reference(diagnostic_ws, min_col=5, min_row=4, max_row=bottleneck_end_row)
        bottleneck_labels = Reference(diagnostic_ws, min_col=4, min_row=5, max_row=bottleneck_end_row)
        bottleneck_chart.add_data(bottleneck_data, titles_from_data=True)
        bottleneck_chart.set_categories(bottleneck_labels)
        bottleneck_chart.height = 7
        bottleneck_chart.width = 11
        diagnostic_ws.add_chart(bottleneck_chart, "M4")

    diagnostic_ws["A24"] = "Bottleneck Details"
    diagnostic_ws["A24"].font = Font(size=12, bold=True)
    program_bottleneck_rows = [
        [item["program"], item["count"], item["avg_age_days"]]
        for item in diagnostic["bottleneck_by_program"]
    ]
    program_bottleneck_end = _write_table(
        diagnostic_ws,
        25,
        1,
        ["Program", "Requests", "Avg Delay (days)"],
        program_bottleneck_rows,
    )
    cert_bottleneck_rows = [
        [item["certificate_type"], item["count"], item["avg_age_days"]]
        for item in diagnostic["bottleneck_by_certificate"]
    ]
    cert_bottleneck_end = _write_table(
        diagnostic_ws,
        25,
        10,
        ["Certificate Type", "Requests", "Avg Delay (days)"],
        cert_bottleneck_rows,
    )

    if program_bottleneck_rows:
        pb_chart = BarChart()
        pb_chart.title = "Bottlenecked Programs"
        pb_chart.y_axis.title = "Value"
        pb_chart.x_axis.title = "Program"
        pb_data = Reference(diagnostic_ws, min_col=2, min_row=25, max_col=3, max_row=program_bottleneck_end)
        pb_labels = Reference(diagnostic_ws, min_col=1, min_row=26, max_row=program_bottleneck_end)
        pb_chart.add_data(pb_data, titles_from_data=True)
        pb_chart.set_categories(pb_labels)
        pb_chart.height = 8
        pb_chart.width = 14
        diagnostic_ws.add_chart(pb_chart, "D24")

    if cert_bottleneck_rows:
        scatter = ScatterChart()
        scatter.title = "Bottlenecked Certificate Types"
        scatter.x_axis.title = "Requests"
        scatter.y_axis.title = "Avg Delay (days)"
        xvalues = Reference(diagnostic_ws, min_col=11, min_row=26, max_row=cert_bottleneck_end)
        yvalues = Reference(diagnostic_ws, min_col=12, min_row=26, max_row=cert_bottleneck_end)
        series = Series(yvalues, xvalues, title="Certificate Types")
        scatter.series.append(series)
        scatter.height = 8
        scatter.width = 14
        diagnostic_ws.add_chart(scatter, "M24")

    predictive_ws = wb.create_sheet("Predictive")
    _style_sheet_title(
        predictive_ws,
        "Predictive Analytics",
        f"Generated at {summary['generated_at']} | Period: {summary['period']}",
    )
    predictive_ws["A4"] = (
        f"Trend window: {predictive['methodology']['linear_trend_window_days']} days"
    )
    daily_rows = [
        [row["date"], row["count"]]
        for row in descriptive["daily_requests"]
    ]
    daily_end_row = _write_table(
        predictive_ws, 6, 1, ["Date", "Actual Requests"], daily_rows
    )
    forecast_rows = [
        [
            row["date"],
            row["linear"],
            row["moving_avg"],
            row["exp_smoothing"],
            row["blended"],
        ]
        for row in predictive["forecast_next_7_days"]
    ]
    forecast_end_row = _write_table(
        predictive_ws,
        6,
        4,
        ["Date", "Linear", "Moving Avg", "Exp Smoothing", "Blended"],
        forecast_rows,
    )

    combined_start = max(daily_end_row, forecast_end_row) + 4
    predictive_ws.cell(row=combined_start, column=1, value="Date").font = Font(bold=True)
    predictive_ws.cell(row=combined_start, column=2, value="Actual Requests").font = Font(bold=True)
    predictive_ws.cell(row=combined_start, column=3, value="Forecast (Blended)").font = Font(bold=True)
    row_cursor = combined_start + 1
    for row in descriptive["daily_requests"]:
        predictive_ws.cell(row=row_cursor, column=1, value=row["date"])
        predictive_ws.cell(row=row_cursor, column=2, value=row["count"])
        predictive_ws.cell(row=row_cursor, column=3, value=None)
        row_cursor += 1
    for row in predictive["forecast_next_7_days"]:
        predictive_ws.cell(row=row_cursor, column=1, value=row["date"])
        predictive_ws.cell(row=row_cursor, column=2, value=None)
        predictive_ws.cell(row=row_cursor, column=3, value=row["blended"])
        row_cursor += 1

    line_chart = LineChart()
    line_chart.title = "Actual Requests vs Forecast"
    line_chart.y_axis.title = "Requests"
    line_chart.x_axis.title = "Date"
    line_chart.height = 9
    line_chart.width = 18
    actual_data = Reference(predictive_ws, min_col=2, min_row=combined_start, max_row=row_cursor - 1)
    forecast_data = Reference(predictive_ws, min_col=3, min_row=combined_start, max_row=row_cursor - 1)
    dates = Reference(predictive_ws, min_col=1, min_row=combined_start + 1, max_row=row_cursor - 1)
    line_chart.add_data(actual_data, titles_from_data=True)
    line_chart.add_data(forecast_data, titles_from_data=True)
    line_chart.set_categories(dates)
    predictive_ws.add_chart(line_chart, "J6")

    stalled_ws = wb.create_sheet("Stalled Requests")
    _style_sheet_title(
        stalled_ws,
        "Stalled Requests",
        f"Generated at {summary['generated_at']} | Period: {summary['period']}",
    )
    stalled_rows = [
        [
            row["reference_number"],
            row["student_name"],
            row["status"],
            row["age_days"],
        ]
        for row in diagnostic["stalled_requests"]
    ]
    _write_table(
        stalled_ws,
        4,
        1,
        ["Reference Number", "Student Name", "Status", "Age (days)"],
        stalled_rows,
    )
    for ws in (descriptive_ws, diagnostic_ws, predictive_ws, stalled_ws):
        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["C"].width = 18
        ws.column_dimensions["D"].width = 24
        ws.column_dimensions["E"].width = 18
        ws.column_dimensions["F"].width = 18
        ws.column_dimensions["G"].width = 24
        ws.column_dimensions["H"].width = 18
        ws.column_dimensions["I"].width = 18
        ws.column_dimensions["J"].width = 24
        ws.column_dimensions["K"].width = 18
        ws.column_dimensions["L"].width = 18

    return wb


def _build_requests_workbook(requests: list):
    wb = Workbook()
    ws = wb.active
    ws.title = "All Requests"

    headers = [
        "Reference Number",
        "Certificate Type",
        "Student Name",
        "Program",
        "Major",
        "Status",
        "Requestor Name",
        "Requestor Email",
        "Created At",
        "Updated At",
    ]
    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True)

    for request in requests:
        ws.append(
            [
                request.reference_number,
                request.certificate_type_name,
                request.student_name,
                request.program,
                request.major,
                request.status.value
                if hasattr(request.status, "value")
                else str(request.status),
                request.requestor_name,
                request.requestor_email,
                _excel_safe_datetime(request.created_at),
                _excel_safe_datetime(request.updated_at),
            ]
        )

    ws.freeze_panes = "A2"

    for column_cells in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(
            max(max_length + 2, 14),
            36,
        )

    return wb


def _filter_request_export_rows(
    requests: list,
    now: datetime,
    search: Optional[str] = None,
    status_filter: Optional[str] = None,
    date_days: Optional[str] = None,
):
    filtered_requests = requests

    if search:
        query = search.strip().lower()
        if query:
            filtered_requests = [
                r
                for r in filtered_requests
                if query
                in " ".join(
                    [
                        str(r.reference_number or ""),
                        str(r.student_name or ""),
                        str(r.certificate_type_name or ""),
                        str(r.program or ""),
                        str(r.sr_code or ""),
                        str(r.requestor_name or ""),
                    ]
                ).lower()
            ]

    if status_filter and status_filter != "ALL":
        filtered_requests = [
            r
            for r in filtered_requests
            if (r.status.value if hasattr(r.status, "value") else str(r.status))
            == status_filter
        ]

    if date_days not in {None, "", "null"}:
        try:
            days = int(date_days)
            start = now.date() - timedelta(days=days)
            filtered_requests = [
                r
                for r in filtered_requests
                if _safe_date(r.created_at) and _safe_date(r.created_at) >= start
            ]
        except ValueError:
            pass

    return filtered_requests


@router.get("/summary")
def get_reports_summary(
    period: Optional[str] = None,
    db: Session = Depends(get_db),
    _: dict = Depends(require_permissions("reports.read")),
):
    request_repo = CertificateRequestRepository(db)
    requests = request_repo.certificates_only().all()
    now = datetime.now()
    return _build_reports_payload(requests, period, now)


@router.get("/export")
def export_reports(
    report_type: str = "summary",
    period: Optional[str] = None,
    search: Optional[str] = None,
    status_filter: Optional[str] = None,
    date_days: Optional[str] = None,
    db: Session = Depends(get_db),
    _: dict = Depends(require_permissions("reports.read")),
):
    request_repo = CertificateRequestRepository(db)
    requests = request_repo.certificates_only().all()
    now = datetime.now()

    if report_type == "requests":
        filtered_requests = requests

        if search:
            query = search.strip().lower()
            if query:
                filtered_requests = [
                    r
                    for r in filtered_requests
                    if query
                    in " ".join(
                        [
                            str(r.reference_number or ""),
                            str(r.student_name or ""),
                            str(r.certificate_type_name or ""),
                            str(r.program or ""),
                            str(r.sr_code or ""),
                            str(r.requestor_name or ""),
                        ]
                    ).lower()
                ]

        if status_filter and status_filter != "ALL":
            filtered_requests = [
                r
                for r in filtered_requests
                if (r.status.value if hasattr(r.status, "value") else str(r.status))
                == status_filter
            ]

        if date_days not in {None, "", "null"}:
            try:
                days = int(date_days)
                start = now.date() - timedelta(days=days)
                filtered_requests = [
                    r
                    for r in filtered_requests
                    if _safe_date(r.created_at) and _safe_date(r.created_at) >= start
                ]
            except ValueError:
                pass

        workbook = _build_requests_workbook(filtered_requests)
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        filename = f"certify_requests_{now.strftime('%Y%m%d_%H%M%S')}.xlsx"
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"},
        )

    summary = _build_reports_payload(requests, period, now)
    workbook = _build_summary_workbook(summary)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"certify_summary_{now.strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/requests-export")
def export_request_rows(
    search: Optional[str] = None,
    status_filter: Optional[str] = None,
    date_days: Optional[str] = None,
    db: Session = Depends(get_db),
    _: dict = Depends(require_permissions("requests.read")),
):
    request_repo = CertificateRequestRepository(db)
    requests = request_repo.certificates_only().all()
    now = datetime.now()
    filtered_requests = _filter_request_export_rows(
        requests,
        now,
        search=search,
        status_filter=status_filter,
        date_days=date_days,
    )

    workbook = _build_requests_workbook(filtered_requests)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"certify_requests_{now.strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
